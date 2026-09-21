"""app.py - ระบบตรวจข้อสอบ OMR (Streamlit)   รัน: streamlit run app.py"""
import io, os, re, sys, tempfile
from datetime import datetime

import pandas as pd
import streamlit as st
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.formatting.rule import FormulaRule, DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter as L

sys.path.insert(0, os.path.join(os.path.dirname(__file__), "src"))
from omr_reader import read_sheet

N_MAX = 170

# ------------------------------------------------------------ ธีมสี/ฟอนต์ ----
NAVY, NAVY_LIGHT, GREY_ROW = "1F3864", "D9E1F2", "F2F2F2"
GREEN, RED, AMBER = "E2EFDA", "FCE4D6", "FFF2CC"
FONT = "TH Sarabun New"
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# เกณฑ์ระดับความยากของข้อสอบ (สัดส่วนผู้ตอบถูก) - แก้ตัวเลขได้ที่นี่
EASY_MIN, MEDIUM_MIN = 0.80, 0.40


def parse_key(text):
    """รับเฉลยได้ 2 แบบ: '1-A 2-C 3-B ...' หรือ 'ACBDE...' (เรียงตามข้อ) -> list ตัวอักษรพิมพ์ใหญ่"""
    pairs = re.findall(r"(\d+)\s*[-:.,)]?\s*([A-Ea-e])", text)
    if pairs and len(pairs) >= 2:
        d = {int(q): a.upper() for q, a in pairs}
        n = max(d)
        return [d.get(i, "") for i in range(1, n + 1)]
    return [c.upper() for c in re.findall(r"[A-Ea-e]", text)]


def grade(answers, key):
    rows = []
    for q, k in enumerate(key, 1):
        a = answers[q - 1]
        rows.append({"ข้อ": q, "ตอบ": a, "เฉลย": k, "ถูก": int(a == k) if k else None})
    return rows


def _ranked(students, key):
    """เรียงผู้สอบจากคะแนนสูงไปต่ำ (คะแนนเท่ากันเรียงตามชื่อไฟล์)"""
    n = len(key)
    return sorted(students, key=lambda x: (-sum(a == k for a, k in zip(x[3][:n], key)), x[0]))


def group_size(m):
    """จำนวนผู้สอบในกลุ่มสูง/กลุ่มต่ำ = 27% ของผู้สอบ (อย่างน้อย 1 คน)"""
    return max(1, round(0.27 * m))


def difficulty_label(p):
    return "ง่าย" if p >= EASY_MIN else ("ปานกลาง" if p >= MEDIUM_MIN else "ยาก")


def discrimination_label(d):
    if d is None or d != d:
        return "-"
    return "ดีมาก" if d >= 0.4 else ("ดี" if d >= 0.3 else ("พอใช้" if d >= 0.2 else "ควรปรับปรุง"))


def item_table(key, students):
    """วิเคราะห์รายข้อ: การกระจายตัวเลือก, ความยาก, อำนาจจำแนก (กลุ่มสูง-ต่ำ 27%)"""
    n, m = len(key), len(students)
    ranked, k = _ranked(students, key), group_size(len(students))
    top, bot = ranked[:k], ranked[-k:]
    rows = []
    for q in range(n):
        col = [s[3][q] for s in students]
        right = sum(a == key[q] for a in col)
        d = ((sum(s[3][q] == key[q] for s in top) - sum(s[3][q] == key[q] for s in bot)) / k
             if m >= 2 else float("nan"))
        rows.append({"ข้อ": q + 1, "เฉลย": key[q], **{c: col.count(c) for c in "ABCDE"},
                     "ไม่ตอบ": col.count("-"), "ฝนซ้อน": col.count("*"), "ตอบถูก": right,
                     "ร้อยละถูก": round(100 * right / m, 1), "ความยาก": difficulty_label(right / m),
                     "อำนาจจำแนก": round(d, 2), "แปลผล": discrimination_label(d)})
    return pd.DataFrame(rows)


def kr20(key, students):
    """ความเชื่อมั่นของข้อสอบ (KR-20); คืน None ถ้าคำนวณไม่ได้"""
    n, m = len(key), len(students)
    if m < 2 or n < 2:
        return None
    scores = [sum(a == k for a, k in zip(s[3][:n], key)) for s in students]
    mean = sum(scores) / m
    var = sum((x - mean) ** 2 for x in scores) / m
    if var == 0:
        return None
    pq = 0.0
    for q in range(n):
        p = sum(s[3][q] == key[q] for s in students) / m
        pq += p * (1 - p)
    return n / (n - 1) * (1 - pq / var)


# ---------------------------------------------------------------- Excel ----
def _f(bold=False, color="000000", size=14):
    return Font(name=FONT, bold=bold, color=color, size=size)


def _header(ws, row, first_col=1, last_col=None):
    last_col = last_col or ws.max_column
    for c in range(first_col, last_col + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = _f(True, "FFFFFF")
        cell.fill = PatternFill("solid", fgColor=NAVY)
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER
    ws.row_dimensions[row].height = 26


def _body(ws, first_row, last_row, first_col, last_col, left_cols=(), zebra=True):
    for r in range(first_row, last_row + 1):
        for c in range(first_col, last_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.font = _f()
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if c in left_cols else "center",
                                       vertical="center")
            if zebra and (r - first_row) % 2 == 1:
                cell.fill = PatternFill("solid", fgColor=GREY_ROW)


def _title(ws, last_col, lines):
    """หัวรายงาน: lines = [(ข้อความ, ขนาด, ตัวหนา)]"""
    for i, (text, size, bold) in enumerate(lines, 1):
        ws.merge_cells(start_row=i, start_column=1, end_row=i, end_column=last_col)
        c = ws.cell(row=i, column=1, value=text)
        c.font = _f(bold, NAVY if bold else "404040", size)
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[i].height = 28 if bold else 22


def _page(ws, landscape=False, title_rows=None):
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.orientation = "landscape" if landscape else "portrait"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.page_margins.left = ws.page_margins.right = 0.5
    ws.oddFooter.center.text = "หน้า &P / &N"
    if title_rows:
        ws.print_title_rows = title_rows
    ws.sheet_view.showGridLines = False


def build_workbook(key, students, meta=None):
    """students = [(ชื่อไฟล์, subject, exam_id, answers)]  คะแนนคำนวณด้วยสูตรใน Excel
    meta = dict(org=, exam=, subject=) ใช้เป็นหัวรายงาน (ไม่บังคับ)"""
    meta = meta or {}
    n = len(key)
    # เรียงจากคะแนนสูงไปต่ำ (คะแนนเท่ากันเรียงตามชื่อไฟล์) ให้ตารางเป็นลำดับอันดับ
    students = sorted(students, key=lambda x: (-sum(a == k for a, k in zip(x[3][:n], key)), x[0]))
    m = len(students)
    wb = Workbook()

    # ---------------- Answers (ข้อมูลดิบ + เฉลย; สีถูก/ผิดด้วย conditional formatting) ----
    a = wb.active
    a.title = "Answers"
    a.append(["ไฟล์", "รหัสวิชา", "เลขประจำตัวสอบ"] + [f"ข้อ {i}" for i in range(1, n + 1)])
    a.append(["เฉลย", "", ""] + key)
    for name, subj, eid, ans in students:
        a.append([name, subj, eid] + ans[:n])
    last_col, last_row = 3 + n, 2 + m
    _header(a, 1, 1, last_col)
    _body(a, 2, last_row, 1, last_col, left_cols=(1,), zebra=False)
    for c in range(1, last_col + 1):                      # แถวเฉลย
        cell = a.cell(row=2, column=c)
        cell.font = _f(True, "1F3864")
        cell.fill = PatternFill("solid", fgColor=NAVY_LIGHT)
    fq, lq = L(4), L(last_col)
    if m:
        rng = f"{fq}3:{lq}{last_row}"
        a.conditional_formatting.add(rng, FormulaRule(
            formula=[f"{fq}3={fq}$2"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
        a.conditional_formatting.add(rng, FormulaRule(
            formula=[f'OR({fq}3="-",{fq}3="*")'],
            fill=PatternFill("solid", bgColor=AMBER, fgColor=AMBER)))
        a.conditional_formatting.add(rng, FormulaRule(
            formula=[f'{fq}3<>{fq}$2'], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
    a.freeze_panes = "D3"
    a.column_dimensions["A"].width = 30
    a.column_dimensions["B"].width = 13
    a.column_dimensions["C"].width = 18
    for c in range(4, last_col + 1):
        a.column_dimensions[L(c)].width = 6.5
    _page(a, landscape=True, title_rows="1:2")

    # ---------------- Summary ----------------
    s = wb.create_sheet("Summary", 0)
    NC = 9                               # จำนวนคอลัมน์ของตาราง
    H = 14                               # แถวหัวตาราง
    d0, d1 = H + 1, H + m                # แถวข้อมูลแรก/สุดท้าย
    now = datetime.now().strftime("%d/%m/%Y %H:%M")
    lines = []
    if meta.get("org"):
        lines.append((meta["org"], 16, True))
    lines.append(("รายงานผลการตรวจข้อสอบ", 18, True))
    detail = " | ".join(x for x in (meta.get("subject"), meta.get("exam")) if x)
    if detail:
        lines.append((detail, 14, False))
    lines.append((f"ตรวจเมื่อ {now}  |  จำนวนข้อ {n} ข้อ", 12, False))
    while len(lines) < 4:
        lines.insert(0, ("", 12, False))
    _title(s, NC, lines[:4])

    # --- สรุปสถิติ ---
    s.merge_cells("A5:E5")
    s["A5"] = "สรุปสถิติ"
    _header(s, 5, 1, 5)
    stats = [
        ("จำนวนผู้เข้าสอบ (คน)", f"=COUNTA(B{d0}:B{d1})", "0"),
        ("คะแนนเฉลี่ย", f"=AVERAGE(E{d0}:E{d1})", "0.00"),
        ("คะแนนสูงสุด", f"=MAX(E{d0}:E{d1})", "0"),
        ("คะแนนต่ำสุด", f"=MIN(E{d0}:E{d1})", "0"),
        ("คะแนนมัธยฐาน", f"=MEDIAN(E{d0}:E{d1})", "0.0"),
        ("ส่วนเบี่ยงเบนมาตรฐาน (S.D.)", f'=IFERROR(STDEV(E{d0}:E{d1}),"-")', "0.00"),
        ("ความเชื่อมั่นของข้อสอบ (KR-20)",
         f'=IFERROR(COUNT(ItemAnalysis!A2:A{n + 1})/(COUNT(ItemAnalysis!A2:A{n + 1})-1)'
         f'*(1-SUM(ItemAnalysis!O2:O{n + 1})/VARP(E{d0}:E{d1})),"-")', "0.00"),
    ]
    for i, (label, formula, fmt) in enumerate(stats):
        r = 6 + i
        s.merge_cells(start_row=r, start_column=1, end_row=r, end_column=4)
        s.cell(row=r, column=1, value=label)
        s.cell(row=r, column=5, value=formula).number_format = fmt
        for c in range(1, 6):
            cell = s.cell(row=r, column=c)
            cell.font = _f(c == 5)
            cell.border = BORDER
            cell.alignment = Alignment(horizontal="left" if c == 1 else "center", indent=1 if c == 1 else 0)
            if c == 5:
                cell.fill = PatternFill("solid", fgColor=NAVY_LIGHT)

    # --- ตารางคะแนนรายคน เรียงตามอันดับ ---
    heads = ["อันดับ", "ไฟล์", "รหัสวิชา", "เลขประจำตัวสอบ", "คะแนน", "เต็ม", "ร้อยละ",
             "ไม่ได้ตอบ", "ฝนซ้อน"]
    for c, h in enumerate(heads, 1):
        s.cell(row=H, column=c, value=h)
    _header(s, H, 1, NC)
    for i in range(m):
        r, ar = H + 1 + i, i + 3
        rng = f"Answers!{fq}{ar}:{lq}{ar}"
        key_rng = f"Answers!${fq}$2:${lq}$2"
        row = [f"=RANK(E{r},$E${d0}:$E${d1})", f"=Answers!A{ar}", f"=Answers!B{ar}", f"=Answers!C{ar}",
               f"=SUMPRODUCT(--({rng}={key_rng}))", n, f"=E{r}/F{r}",
               f'=COUNTIF({rng},"-")', f'=COUNTIF({rng},"~*")']
        for c, v in enumerate(row, 1):
            s.cell(row=r, column=c, value=v)
    _body(s, d0, d1, 1, NC, left_cols=(2,))
    for r in range(d0, d1 + 1):
        s[f"G{r}"].number_format = "0.0%"
        s[f"A{r}"].font = _f(True, NAVY)
    if m:
        s.conditional_formatting.add(f"C{d0}:D{d1}", FormulaRule(
            formula=[f'ISNUMBER(SEARCH("~?",C{d0}))'],
            fill=PatternFill("solid", bgColor="F8CBAD", fgColor="F8CBAD")))
        s.conditional_formatting.add(f"A{d0}:{L(NC)}{d1}", FormulaRule(
            formula=[f"$A{d0}<=3"], fill=PatternFill("solid", bgColor="FFE699", fgColor="FFE699")))
        s.conditional_formatting.add(f"G{d0}:G{d1}", DataBarRule(
            start_type="num", start_value=0, end_type="num", end_value=1, color="8EA9DB"))
    note = d1 + 2
    s.merge_cells(start_row=note, start_column=1, end_row=note, end_column=NC)
    s.cell(row=note, column=1,
           value="หมายเหตุ: แถวสีทอง = 3 อันดับแรก (คะแนนเท่ากันได้อันดับเดียวกัน) | ช่องรหัสสีส้ม = อ่านรหัสไม่ได้ (?) "
                 "ควรตรวจกับกระดาษคำตอบจริง | ไม่ได้ตอบ = ไม่ฝนข้อนั้น | ฝนซ้อน = ฝนมากกว่า 1 ตัวเลือก (นับเป็นผิด)"
           ).font = _f(size=12, color="595959")
    s.cell(row=note, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    s.row_dimensions[note].height = 44
    for col, w in zip("ABCDEFGHI", (9, 30, 13, 18, 11, 9, 13, 12, 11)):
        s.column_dimensions[col].width = w
    s.freeze_panes = f"A{H + 1}"
    _page(s, title_rows=f"{H}:{H}")

    # ---------------- ItemAnalysis ----------------
    it = wb.create_sheet("ItemAnalysis")
    heads = ["ข้อ", "เฉลย", "A", "B", "C", "D", "E", "ไม่ตอบ", "ฝนซ้อน", "ตอบถูก",
             "ร้อยละถูก", "ความยาก", "อำนาจจำแนก", "แปลผล", "p(1-p)"]
    it.append(heads)
    _header(it, 1, 1, 15)
    nk = n + 3                                   # แถวที่เก็บจำนวนผู้สอบในกลุ่มสูง/ต่ำ
    KC = f"$C${nk}"
    for q in range(1, n + 1):
        col, r = L(3 + q), q + 1
        rng = f"Answers!{col}$3:{col}${last_row}"
        keyc = f"Answers!{col}$2"
        top = f"Answers!{col}$3:INDEX(Answers!{col}:{col},2+{KC})"
        bot = f"INDEX(Answers!{col}:{col},{last_row}-{KC}+1):Answers!{col}${last_row}"
        it.append([
            q, f"={keyc}",
            *[f"=COUNTIF({rng},{L(3 + j)}$1)" for j in range(5)],
            f'=COUNTIF({rng},"-")', f'=COUNTIF({rng},"~*")',
            f"=SUMPRODUCT(--({rng}={keyc}))", f"=J{r}/{m}",
            f'=IF(K{r}>={EASY_MIN},"ง่าย",IF(K{r}>={MEDIUM_MIN},"ปานกลาง","ยาก"))',
            f'=IFERROR((SUMPRODUCT(--({top}={keyc}))-SUMPRODUCT(--({bot}={keyc})))/{KC},"-")',
            f'=IF(ISNUMBER(M{r}),IF(M{r}>=0.4,"ดีมาก",IF(M{r}>=0.3,"ดี",IF(M{r}>=0.2,"พอใช้","ควรปรับปรุง"))),"-")',
            f"=K{r}*(1-K{r})"])
    _body(it, 2, n + 1, 1, 15)
    for r in range(2, n + 2):
        it[f"K{r}"].number_format = "0.0%"
        it[f"M{r}"].number_format = "0.00"
        it[f"O{r}"].number_format = "0.000"
    last_it = n + 1
    it.conditional_formatting.add(f"C2:G{last_it}", FormulaRule(       # ตัวเลือกที่เป็นเฉลย
        formula=["C$1=$B2"], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN),
        font=Font(name=FONT, bold=True, color="375623")))
    it.conditional_formatting.add(f"K2:K{last_it}", DataBarRule(
        start_type="num", start_value=0, end_type="num", end_value=1, color="8EA9DB"))
    it.conditional_formatting.add(f"N2:N{last_it}", FormulaRule(
        formula=['N2="ควรปรับปรุง"'], fill=PatternFill("solid", bgColor=RED, fgColor=RED)))
    it.conditional_formatting.add(f"N2:N{last_it}", FormulaRule(
        formula=['N2="ดีมาก"'], fill=PatternFill("solid", bgColor=GREEN, fgColor=GREEN)))
    it.conditional_formatting.add(f"L2:L{last_it}", FormulaRule(
        formula=['L2="ยาก"'], fill=PatternFill("solid", bgColor=AMBER, fgColor=AMBER)))
    it.merge_cells(start_row=nk, start_column=1, end_row=nk, end_column=2)
    it.cell(row=nk, column=1, value="กลุ่มสูง/ต่ำ (คน)")
    it.cell(row=nk, column=3, value=group_size(m))
    for c in (1, 2, 3):
        cell = it.cell(row=nk, column=c)
        cell.font = _f(True, "0000FF" if c == 3 else "000000")
        cell.border = BORDER
        cell.alignment = Alignment(horizontal="center")
    it.merge_cells(start_row=nk + 1, start_column=1, end_row=nk + 3, end_column=15)
    it.cell(row=nk + 1, column=1, value=(
        f"หมายเหตุ: A-E = จำนวนผู้เลือกแต่ละตัวเลือก (สีเขียว = เฉลย) | ความยาก: ง่าย ≥ {EASY_MIN:.0%}, "
        f"ปานกลาง ≥ {MEDIUM_MIN:.0%}, ยาก < {MEDIUM_MIN:.0%} | อำนาจจำแนก = (ผู้ตอบถูกกลุ่มสูง - กลุ่มต่ำ) / "
        "จำนวนกลุ่ม โดยกลุ่มสูง/ต่ำคือผู้สอบ 27% แรก/ท้ายตามอันดับ (แก้จำนวนคนได้ที่ช่องสีน้ำเงิน) | "
        "เกณฑ์แปลผล: ≥ 0.40 ดีมาก, 0.30-0.39 ดี, 0.20-0.29 พอใช้, < 0.20 ควรปรับปรุง | "
        "ผู้สอบในชีต Answers เรียงตามคะแนนขณะออกรายงาน หากแก้เฉลยภายหลัง ควรออกรายงานใหม่")
    ).font = _f(size=12, color="595959")
    it.cell(row=nk + 1, column=1).alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEFGHIJKLMNO", (7, 8, 7, 7, 7, 7, 7, 9, 10, 10, 13, 12, 14, 15, 10)):
        it.column_dimensions[col].width = w
    it.freeze_panes = "C2"
    ch = BarChart()
    ch.type, ch.title = "col", "ร้อยละของผู้ตอบถูกรายข้อ"
    ch.add_data(Reference(it, min_col=11, min_row=1, max_row=last_it), titles_from_data=True)
    ch.set_categories(Reference(it, min_col=1, min_row=2, max_row=last_it))
    ch.y_axis.scaling.min, ch.y_axis.scaling.max = 0, 1
    ch.y_axis.number_format = "0%"
    ch.x_axis.delete = ch.y_axis.delete = False
    ch.legend = None
    ch.series[0].graphicalProperties.solidFill = NAVY
    ch.width, ch.height = 30, 10
    it.add_chart(ch, "Q2")
    _page(it, landscape=True, title_rows="1:1")

    # ---------------- Distribution (การกระจายคะแนน) ----------------
    ds = wb.create_sheet("Distribution")
    ds.append(["ช่วงคะแนน (ร้อยละ)", "จาก (≥)", "ถึง (<)", "จำนวนคน", "สัดส่วน"])
    _header(ds, 1, 1, 5)
    for i in range(10):
        r = i + 2
        lo, hi = i / 10, ((i + 1) / 10 if i < 9 else 1.0001)
        label = f"{i * 10}-{(i + 1) * 10 - 1}%" if i < 9 else "90-100%"
        g = f"Summary!$G${d0}:$G${d1}"
        ds.append([label, lo, hi, f'=COUNTIFS({g},">="&B{r},{g},"<"&C{r})',
                   f"=IFERROR(D{r}/SUM($D$2:$D$11),0)"])
    ds.append(["รวม", "", "", "=SUM(D2:D11)", "=SUM(E2:E11)"])
    _body(ds, 2, 12, 1, 5)
    for r in range(2, 13):
        ds[f"B{r}"].number_format = ds[f"C{r}"].number_format = "0%"
        ds[f"E{r}"].number_format = "0.0%"
    for c in range(1, 6):
        ds.cell(row=12, column=c).font = _f(True)
        ds.cell(row=12, column=c).fill = PatternFill("solid", fgColor=NAVY_LIGHT)
    for col, w in zip("ABCDE", (22, 10, 10, 12, 12)):
        ds.column_dimensions[col].width = w
    ch2 = BarChart()
    ch2.type, ch2.title = "col", "การกระจายคะแนน (จำนวนคน)"
    ch2.add_data(Reference(ds, min_col=4, min_row=1, max_row=11), titles_from_data=True)
    ch2.set_categories(Reference(ds, min_col=1, min_row=2, max_row=11))
    ch2.x_axis.delete = ch2.y_axis.delete = False
    ch2.legend = None
    ch2.series[0].graphicalProperties.solidFill = NAVY
    ch2.width, ch2.height = 18, 9
    ds.add_chart(ch2, "G2")
    _page(ds, landscape=True)

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


# ------------------------------------------------------------------ UI ----
st.set_page_config(page_title="ระบบตรวจข้อสอบ OMR", layout="wide")

st.markdown(f"""
<style>
@import url('https://fonts.googleapis.com/css2?family=Sarabun:wght@400;600;700&display=swap');
html, body, [class*="css"], .stApp {{ font-family: 'Sarabun', 'TH Sarabun New', sans-serif; }}
#MainMenu, footer, [data-testid="stToolbar"] {{ visibility: hidden; }}
.block-container {{ padding-top: 3rem; max-width: 1200px; }}
.banner {{ background: #{NAVY}; color: #fff; padding: 1.1rem 1.6rem; border-radius: 6px; margin-bottom: 1.2rem; }}
.banner h1 {{ color: #fff; margin: 0; font-size: 1.7rem; font-weight: 700; padding: 0; }}
.banner p {{ color: #D9E1F2; margin: .2rem 0 0; font-size: 1rem; }}
.step {{ border-left: 4px solid #{NAVY}; padding-left: .7rem; margin: .4rem 0 .6rem;
        font-weight: 700; font-size: 1.15rem; color: #{NAVY}; }}
[data-testid="stMetric"] {{ background: #F3F5F9; border: 1px solid #D9E1F2; border-radius: 6px; padding: .7rem 1rem; }}
.rank-card {{ border: 1px solid #D9E1F2; border-top: 4px solid #C9A227; border-radius: 6px; padding: .8rem 1rem; background: #FFFDF5; }}
.rank-no {{ color: #{NAVY}; font-weight: 700; }}
.rank-score {{ font-size: 2rem; font-weight: 700; color: #{NAVY}; line-height: 1.2; }}
.rank-score span {{ font-size: 1rem; color: #7F7F7F; font-weight: 400; }}
.rank-name {{ font-weight: 600; word-break: break-all; }}
.rank-sub {{ color: #595959; font-size: .9rem; }}
[data-testid="stMetricLabel"] {{ color: #595959; }}
</style>
<div class="banner">
  <h1>ระบบตรวจข้อสอบ OMR</h1>
  <p>กระดาษคำตอบแบบ R1101 (สูงสุด {N_MAX} ข้อ, 5 ตัวเลือก) &nbsp;|&nbsp; ส่งออกรายงานเป็นไฟล์ Excel</p>
</div>
""", unsafe_allow_html=True)


def step(text):
    st.markdown(f'<div class="step">{text}</div>', unsafe_allow_html=True)


c1, c2 = st.columns(2, gap="large")
with c1:
    step("ขั้นตอนที่ 1  ข้อมูลการสอบและเฉลย")
    with st.expander("ข้อมูลหัวรายงาน (ไม่บังคับ)"):
        org = st.text_input("หน่วยงาน / สถาบัน")
        subj_name = st.text_input("ชื่อรายวิชา")
        exam_name = st.text_input("การสอบ (เช่น สอบกลางภาค ภาคเรียนที่ 1/2569)")
    key_text = st.text_area(
        "เฉลย  (เช่น ABCDE... เรียงตามข้อ หรือ 1-A 2-C 3-B ...)",
        height=140, placeholder="ABCDEABCDE...")
    key = parse_key(key_text)
    if key:
        st.caption(f"อ่านเฉลยได้ {len(key)} ข้อ")
    else:
        st.caption("ยังไม่ได้ใส่เฉลย")
with c2:
    step("ขั้นตอนที่ 2  อัปโหลดกระดาษคำตอบ")
    files = st.file_uploader("ไฟล์ PNG / JPG / PDF (เลือกได้หลายไฟล์)",
                             type=["png", "jpg", "jpeg", "pdf"], accept_multiple_files=True)
    if files:
        st.caption(f"เลือกแล้ว {len(files)} ไฟล์")

step("ขั้นตอนที่ 3  ตรวจและดูผล")
if st.button("เริ่มตรวจข้อสอบ", type="primary", disabled=not (key and files)):
    if len(key) > N_MAX:
        st.error(f"เฉลยยาวเกิน {N_MAX} ข้อ")
        st.stop()
    students, errors = [], []
    bar = st.progress(0.0, text="กำลังตรวจ...")
    for i, f in enumerate(files):
        suffix = os.path.splitext(f.name)[1]
        with tempfile.NamedTemporaryFile(delete=False, suffix=suffix) as tmp:
            tmp.write(f.getvalue())
            path = tmp.name
        try:
            r = read_sheet(path)
            students.append((f.name, r["subject_code"], r["exam_id"], r["answers"]))
        except Exception as e:
            errors.append(f"{f.name}: {e}")
        finally:
            os.unlink(path)
        bar.progress((i + 1) / len(files), text=f"กำลังตรวจ {i + 1}/{len(files)}")
    bar.empty()
    meta = {"org": org, "subject": subj_name, "exam": exam_name}
    # เก็บผลไว้ใน session_state เพื่อไม่ให้ผลหายเมื่อกดดาวน์โหลด
    st.session_state["result"] = {
        "key": key, "students": students, "errors": errors,
        "xlsx": build_workbook(key, students, meta) if students else None,
    }

res = st.session_state.get("result")
if res:
    for e in res["errors"]:
        st.warning(f"อ่านไฟล์ไม่สำเร็จ - {e}")
    students, key = res["students"], res["key"]
    if students:
        n = len(key)
        rows = []
        for name, subj, eid, ans in students:
            g = grade(ans, key)
            score = sum(x["ถูก"] or 0 for x in g)
            pct = round(100 * score / n, 1)
            rows.append({"ไฟล์": name, "รหัสวิชา": subj, "เลขประจำตัวสอบ": eid,
                         "คะแนน": score, "เต็ม": n, "ร้อยละ": pct,
                         "ไม่ได้ตอบ": ans[:n].count("-"), "ฝนซ้อน": ans[:n].count("*")})
        df = pd.DataFrame(rows)
        df.insert(0, "อันดับ", df["คะแนน"].rank(method="min", ascending=False).astype(int))
        df = df.sort_values(["อันดับ", "ไฟล์"]).reset_index(drop=True)

        bad = df[df["รหัสวิชา"].str.contains(r"\?") | df["เลขประจำตัวสอบ"].str.contains(r"\?")]
        if len(bad):
            st.warning(f"อ่านรหัสไม่ได้ {len(bad)} ฉบับ (แสดงเป็น ?) - "
                       "ควรตรวจกับกระดาษคำตอบจริง: " + ", ".join(bad["ไฟล์"].head(5)))

        # ตรวจคุณภาพการอ่าน: ฉบับที่ไม่ตอบ/ฝนซ้อนมากผิดปกติ (>10% ของข้อ) มักเกิดจากการสแกนหรืออ่านผิด
        odd = df[(df["ไม่ได้ตอบ"] + df["ฝนซ้อน"]) > 0.10 * n]
        if len(odd):
            st.warning(f"พบ {len(odd)} ฉบับที่ไม่ได้ตอบ/ฝนซ้อนเกิน 10% ของข้อ ควรตรวจกับกระดาษจริง: "
                       + ", ".join(odd["ไฟล์"].head(5)) + (" ..." if len(odd) > 5 else ""))

        kr = kr20(key, students)
        m = st.columns(6)
        m[0].metric("ผู้เข้าสอบ (คน)", len(df))
        m[1].metric("คะแนนเฉลี่ย", f"{df['คะแนน'].mean():.2f}")
        m[2].metric("สูงสุด", int(df["คะแนน"].max()))
        m[3].metric("ต่ำสุด", int(df["คะแนน"].min()))
        m[4].metric("S.D.", f"{df['คะแนน'].std():.2f}" if len(df) > 1 else "-")
        m[5].metric("ความเชื่อมั่น (KR-20)", f"{kr:.2f}" if kr is not None else "-",
                    help="ยิ่งใกล้ 1 ยิ่งเชื่อถือได้ (โดยทั่วไป ≥ 0.70 ถือว่ายอมรับได้) ต้องมีผู้สอบตั้งแต่ 2 คน")

        # 3 อันดับแรก
        top = df[df["อันดับ"] <= 3]
        st.markdown('<div class="step">อันดับสูงสุด</div>', unsafe_allow_html=True)
        tcols = st.columns(min(3, len(top)))
        for col, (_, r) in zip(tcols, top.head(3).iterrows()):
            col.markdown(
                f'<div class="rank-card"><div class="rank-no">อันดับ {r["อันดับ"]}</div>'
                f'<div class="rank-score">{r["คะแนน"]}<span>/{r["เต็ม"]}</span></div>'
                f'<div class="rank-name">{r["ไฟล์"]}</div>'
                f'<div class="rank-sub">เลขประจำตัวสอบ {r["เลขประจำตัวสอบ"]}</div></div>',
                unsafe_allow_html=True)
        if len(top) > 3:
            st.caption(f"มีผู้ได้อันดับ 1-3 รวม {len(top)} คน (คะแนนเท่ากัน) แสดง 3 คนแรก - ดูทั้งหมดในตาราง")

        tab1, tab2, tab3, tab4 = st.tabs(
            ["ผลรายบุคคล", "ตรวจรายฉบับ", "วิเคราะห์รายข้อ", "การกระจายคะแนน"])
        with tab1:
            def hl(row):
                return ["background-color: #FFF2CC; font-weight: 600" if row["อันดับ"] <= 3 else ""] * len(row)
            st.dataframe(df.style.apply(hl, axis=1), width="stretch", hide_index=True,
                         column_config={"ร้อยละ": st.column_config.ProgressColumn(
                             "ร้อยละ", min_value=0, max_value=100, format="%.1f")})
        with tab2:
            st.caption("ตรวจว่าระบบอ่านคำตอบของแต่ละฉบับถูกต้องหรือไม่ (เทียบกับเฉลย)")
            by_name = {s[0]: s for s in students}
            idx = st.selectbox(
                "เลือกผู้สอบ", range(len(df)),
                format_func=lambda i: f"อันดับ {df.loc[i, 'อันดับ']} | {df.loc[i, 'ไฟล์']} | "
                                      f"เลขประจำตัว {df.loc[i, 'เลขประจำตัวสอบ']} | {df.loc[i, 'คะแนน']}/{n}")
            ans = by_name[df.loc[idx, "ไฟล์"]][3]
            det = pd.DataFrame({
                "ข้อ": range(1, n + 1), "ตอบ": ans[:n], "เฉลย": key,
                "ผล": ["ไม่ตอบ" if a == "-" else "ฝนซ้อน" if a == "*" else "ถูก" if a == k else "ผิด"
                       for a, k in zip(ans[:n], key)]})
            only_bad = st.checkbox("แสดงเฉพาะข้อที่ไม่ถูก")
            if only_bad:
                det = det[det["ผล"] != "ถูก"]
            color = {"ถูก": "#E2EFDA", "ผิด": "#FCE4D6", "ไม่ตอบ": "#FFF2CC", "ฝนซ้อน": "#FFF2CC"}
            st.dataframe(det.style.apply(lambda r: [f"background-color: {color[r['ผล']]}"] * len(r), axis=1),
                         width="stretch", hide_index=True, height=420)
        with tab3:
            items = item_table(key, students)
            weak = items[items["แปลผล"] == "ควรปรับปรุง"]
            hard = items[items["ความยาก"] == "ยาก"]
            c1_, c2_, c3_ = st.columns(3)
            c1_.metric("ข้อที่ง่าย (≥ 80%)", int((items["ความยาก"] == "ง่าย").sum()))
            c2_.metric("ข้อที่ยาก (< 40%)", len(hard))
            c3_.metric("ข้อที่ควรปรับปรุง", len(weak),
                       help="อำนาจจำแนก < 0.20 (กลุ่มสูง/ต่ำ 27%) แยกผู้เก่ง-อ่อนได้น้อย ควรตรวจคำถามและตัวเลือก")
            st.dataframe(
                items.style.apply(lambda r: ["background-color: #FCE4D6" if r["แปลผล"] == "ควรปรับปรุง" else ""] * len(r), axis=1),
                width="stretch", hide_index=True, height=420,
                column_config={"ร้อยละถูก": st.column_config.ProgressColumn(
                    "ร้อยละถูก", min_value=0, max_value=100, format="%.1f")})
            st.caption("A-E = จำนวนผู้เลือกแต่ละตัวเลือก | แถวสีแดง = อำนาจจำแนกต่ำ")
            st.bar_chart(items.set_index("ข้อ")[["ร้อยละถูก"]], color="#" + NAVY)
        with tab4:
            bins = [f"{i * 10}-{i * 10 + 9}%" if i < 9 else "90-100%" for i in range(10)]
            b_idx = (df["ร้อยละ"] // 10).clip(upper=9).astype(int)
            cnt = b_idx.value_counts().reindex(range(10), fill_value=0)
            st.caption("จำนวนผู้สอบในแต่ละช่วงคะแนน")
            st.bar_chart(pd.DataFrame({"ช่วงคะแนน": bins, "จำนวนคน": cnt.values}).set_index("ช่วงคะแนน"),
                         color="#" + NAVY)

        st.download_button("ดาวน์โหลดรายงาน Excel", res["xlsx"],
                           file_name=f"exam_result_{datetime.now():%Y%m%d_%H%M}.xlsx",
                           mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                           type="primary")
