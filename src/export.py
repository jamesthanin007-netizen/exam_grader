"""export.py - อ่านกระดาษ OMR แล้วส่งออก Excel (คะแนนคำนวณด้วยสูตรใน Excel)"""
import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from omr_reader import read_sheet

F = "Arial"
YELLOW = PatternFill("solid", fgColor="FFFF00")
HEAD = PatternFill("solid", fgColor="7030A0")

def export(sheet_path, out_path, key=None):
    r = read_sheet(sheet_path)
    key = key or {}
    wb = Workbook()
    ws = wb.active; ws.title = "Answers"

    hdr = ["ข้อ", "คำตอบที่อ่านได้", "เฉลย", "ถูก(1)/ผิด(0)"]
    ws.append(hdr)
    for c in ws[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HEAD
        c.alignment = Alignment(horizontal="center")
    for q in range(1, 171):
        row = q + 1
        ws.cell(row, 1, q)
        ws.cell(row, 2, r["answers"][q - 1])
        k = ws.cell(row, 3, key.get(q))
        k.fill = YELLOW; k.font = Font(name=F, color="0000FF")
        ws.cell(row, 4, f'=IF(C{row}="","",IF(B{row}=C{row},1,0))')
        for col in (1, 2, 4):
            ws.cell(row, col).font = Font(name=F)
        for col in range(1, 5):
            ws.cell(row, col).alignment = Alignment(horizontal="center")
    for col, w in zip("ABCD", (8, 18, 10, 16)):
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "A2"

    s = wb.create_sheet("Summary", 0)
    rows = [
        ("รหัสวิชา", r["subject_code"]),
        ("เลขประจำตัวสอบ", r["exam_id"]),
        ("จำนวนข้อที่มีเฉลย", "=COUNTA(Answers!C2:C171)"),
        ("คะแนนที่ได้", "=SUM(Answers!D2:D171)"),
        ("เปอร์เซ็นต์", '=IF(B4=0,0,B5/B4)'),
        ("ข้อที่ไม่ได้ตอบ (-)", '=COUNTIF(Answers!B2:B171,"-")'),
        ("ข้อที่ฝนซ้อน/กำกวม (*)", '=COUNTIF(Answers!B2:B171,"~*")'),
    ]
    s.append(["รายการ", "ค่า"])
    for c in s[1]:
        c.font = Font(name=F, bold=True, color="FFFFFF"); c.fill = HEAD
    for label, val in rows:
        s.append([label, val])
    s["B6"].number_format = "0.0%"
    for row in s.iter_rows(min_row=2):
        for c in row: c.font = Font(name=F)
    s["A10"] = "หมายเหตุ"; s["A10"].font = Font(name=F, bold=True)
    s["A11"] = "- ช่องสีเหลืองในชีต Answers (คอลัมน์ C) คือเฉลย ให้แก้เป็นเฉลยจริง"
    s["A12"] = "- เฉลยที่ใส่ไว้ตอนนี้เป็นค่าตัวอย่างสำหรับทดสอบเท่านั้น"
    s["A13"] = "- '-' = ไม่ได้ฝน, '*' = ฝนหลายช่อง/กำกวม"
    for a in ("A11", "A12", "A13"): s[a].font = Font(name=F)
    s.column_dimensions["A"].width = 28; s.column_dimensions["B"].width = 16
    wb.save(out_path)

if __name__ == "__main__":
    demo_key = {q: "ABCDE"[(q - 1) % 5] for q in range(1, 171)}  # เฉลยตัวอย่าง วนซ้ำ A-E
    export(sys.argv[1], sys.argv[2], demo_key)
